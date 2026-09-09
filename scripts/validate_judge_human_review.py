"""
Step 4F: Human Audit Review Validation Script

Validates eval/judge_human_review_ready_for_manual_check.xlsx:
1. Exactly 50 rows in 'Judge-Human Audit'.
2. All message IDs match canonical audit set in exact order.
3. Machine columns unchanged:
   - original_text
   - agent_response
   - llm_judge_* scores
4. Human review columns:
   - human_correctness in {1, 2, 3, 4, 5}
   - human_helpfulness in {1, 2, 3, 4, 5}
   - human_groundedness in {1, 2, 3, 4, 5}
   - human_policy in {1, 2, 3, 4, 5}
   - human_escalation in {1, 2, 3, 4, 5}
   - review_status in {"Reviewed", "Needs Human Review"}
   - human_review_action in {"approved", "edited", None, ""}
5. Reports:
   - 50 total
   - reviewed count
   - pending count
   - approved count
   - edited count
   - machine columns unchanged
"""

import argparse
import csv
import sys
from pathlib import Path
from typing import Dict, Any, List, Optional
import openpyxl

# UTF-8 stdout configuration
if sys.stdout.encoding != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

ROOT_DIR = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = ROOT_DIR / "eval/judge_human_review_ready_for_manual_check.xlsx"
CANONICAL_AUDIT_CSV = ROOT_DIR / "eval/judge_human_audit.csv"

# Column indexes in 'Judge-Human Audit'
COL_MESSAGE_ID = 1
COL_ORIGINAL_TEXT = 2
COL_AGENT_RESPONSE = 3
COL_RETRIEVED_EVIDENCE = 4
COL_AUDIT_FLAG = 5
COL_LLM_CORRECTNESS = 6
COL_ASST_CORRECTNESS = 7
COL_LLM_HELPFULNESS = 8
COL_ASST_HELPFULNESS = 9
COL_LLM_GROUNDEDNESS = 10
COL_ASST_GROUNDEDNESS = 11
COL_LLM_POLICY = 12
COL_ASST_POLICY = 13
COL_LLM_ESCALATION = 14
COL_ASST_ESCALATION = 15
COL_ASSISTANT_REASON = 16
COL_REVIEW_STATUS = 17
COL_HUMAN_CORRECTNESS = 18
COL_HUMAN_HELPFULNESS = 19
COL_HUMAN_GROUNDEDNESS = 20
COL_HUMAN_POLICY = 21
COL_HUMAN_ESCALATION = 22
COL_HUMAN_REVIEW_ACTION = 23
COL_HUMAN_NOTES = 24


def validate_review(workbook_path: Path, require_complete: bool = False) -> bool:
    print("=" * 75)
    print("STEP 4F: VALIDATING HUMAN AUDIT WORKBOOK INTEGRITY")
    print(f"Target: {workbook_path.name}")
    print("=" * 75)

    if not workbook_path.exists():
        print(f"❌ Error: Workbook not found: {workbook_path}")
        return False

    if not CANONICAL_AUDIT_CSV.exists():
        print(f"❌ Error: Canonical audit CSV missing: {CANONICAL_AUDIT_CSV}")
        return False

    # 1. Load canonical audit references
    with open(CANONICAL_AUDIT_CSV, "r", encoding="utf-8") as f:
        canonical_rows = list(csv.DictReader(f))

    if len(canonical_rows) != 50:
        print(f"❌ Error: Expected 50 rows in canonical audit CSV, found {len(canonical_rows)}")
        return False

    canonical_by_id = {str(r["message_id"]): r for r in canonical_rows}
    canonical_mids = [str(r["message_id"]) for r in canonical_rows]

    # 2. Load workbook
    wb = openpyxl.load_workbook(workbook_path, data_only=True)

    if "Judge-Human Audit" not in wb.sheetnames:
        print(f"❌ Error: Sheet 'Judge-Human Audit' not found in {workbook_path.name}.")
        return False

    ws = wb["Judge-Human Audit"]

    # 3. Verify exactly 50 message rows (rows 6 to 55)
    data_start = 6
    data_end = 55

    row_mids = []
    for r in range(data_start, data_end + 1):
        val = ws.cell(row=r, column=COL_MESSAGE_ID).value
        if val is not None and str(val).strip() != "":
            row_mids.append(str(val).strip())

    if len(row_mids) != 50:
        print(f"❌ Integrity Violation: Expected exactly 50 message rows, found {len(row_mids)}.")
        return False
    print("✓ Row Count: Exactly 50 audit rows verified.")

    # 4. Verify message IDs match canonical audit set
    if row_mids != canonical_mids:
        missing = set(canonical_mids) - set(row_mids)
        invalid = set(row_mids) - set(canonical_mids)
        if missing:
            print(f"❌ Missing canonical IDs: {missing}")
        if invalid:
            print(f"❌ Invalid / unexpected IDs: {invalid}")
        return False
    print("✓ ID Conformance: All 50 message IDs match canonical audit set in exact order.")

    # 5. Verify machine columns are NOT overwritten
    machine_mismatches = []
    for r_idx, mid in enumerate(row_mids, start=data_start):
        canon = canonical_by_id[mid]
        # Text
        sheet_text = str(ws.cell(row=r_idx, column=COL_ORIGINAL_TEXT).value or "").strip()
        canon_text = str(canon["original_text"]).strip()
        if sheet_text != canon_text:
            machine_mismatches.append(f"Row {r_idx} (ID {mid}): text modified.")

        # Agent response
        sheet_resp = str(ws.cell(row=r_idx, column=COL_AGENT_RESPONSE).value or "").strip()
        canon_resp = str(canon["agent_response"]).strip()
        if sheet_resp != canon_resp:
            machine_mismatches.append(f"Row {r_idx} (ID {mid}): agent response modified.")

        # LLM judge scores
        judge_cols = [
            (COL_LLM_CORRECTNESS, "llm_judge_correctness"),
            (COL_LLM_HELPFULNESS, "llm_judge_helpfulness"),
            (COL_LLM_GROUNDEDNESS, "llm_judge_groundedness"),
            (COL_LLM_POLICY, "llm_judge_policy"),
            (COL_LLM_ESCALATION, "llm_judge_escalation"),
        ]
        for col_idx, key in judge_cols:
            val = ws.cell(row=r_idx, column=col_idx).value
            if val is None or int(val) != int(canon[key]):
                machine_mismatches.append(f"Row {r_idx} (ID {mid}): {key} altered from {canon[key]} to {val}")

    if machine_mismatches:
        print(f"❌ Error: {len(machine_mismatches)} machine values were overwritten:")
        for m in machine_mismatches[:5]:
            print(f"   - {m}")
        return False
    print("✓ Machine Column Guardrail: All machine/LLM traces and scores remain pristine.")

    # 6. Audit Human Review Fields
    reviewed_count = 0
    pending_count = 0
    approved_count = 0
    edited_count = 0
    invalid_ratings = []
    invalid_statuses = []

    for r_idx, mid in enumerate(row_mids, start=data_start):
        status = str(ws.cell(row=r_idx, column=COL_REVIEW_STATUS).value or "").strip()
        action = str(ws.cell(row=r_idx, column=COL_HUMAN_REVIEW_ACTION).value or "").strip()

        # Check ratings
        ratings = []
        for c in range(COL_HUMAN_CORRECTNESS, COL_HUMAN_ESCALATION + 1):
            val = ws.cell(row=r_idx, column=c).value
            try:
                score = int(val)
                ratings.append(score)
                if score < 1 or score > 5:
                    invalid_ratings.append(f"Row {r_idx} (ID {mid}): Col {c} out-of-range rating '{val}' (must be 1-5)")
            except (ValueError, TypeError):
                invalid_ratings.append(f"Row {r_idx} (ID {mid}): Col {c} non-integer rating '{val}'")

        # Check status and action coherence
        if status == "Reviewed":
            if action not in ["approved", "edited"]:
                invalid_statuses.append(
                    f"Row {r_idx} (ID {mid}): status is 'Reviewed' but action is '{action}' (expected 'approved' or 'edited')"
                )
            else:
                reviewed_count += 1
                if action == "approved":
                    approved_count += 1
                elif action == "edited":
                    edited_count += 1
        elif status == "Needs Human Review":
            if action not in ["", None, "None"]:
                invalid_statuses.append(
                    f"Row {r_idx} (ID {mid}): status is 'Needs Human Review' but action is already set to '{action}'"
                )
            pending_count += 1
        else:
            invalid_statuses.append(
                f"Row {r_idx} (ID {mid}): unexpected review_status '{status}' (expected 'Reviewed' or 'Needs Human Review')"
            )
            pending_count += 1

    # 7. Summary Dashboard Reporting
    print("-" * 75)
    print("HUMAN AUDIT VALIDATION REPORT:")
    print(f"Total Audit Cases:        50")
    print(f"Reviewed Count:           {reviewed_count} / 50 ({reviewed_count/50*100:.1f}%)")
    print(f"Pending Count:            {pending_count} / 50 ({pending_count/50*100:.1f}%)")
    print(f"Approved Count:           {approved_count}")
    print(f"Edited Count:             {edited_count}")
    print(f"Machine Columns Unchanged: VERIFIED (All 50 cases match canonical traces)")
    print("-" * 75)

    if invalid_ratings:
        print(f"❌ Error: {len(invalid_ratings)} invalid rating entries found:")
        for err in invalid_ratings[:5]:
            print(f"   - {err}")
        return False

    if invalid_statuses:
        print(f"❌ Error: {len(invalid_statuses)} invalid review status/action entries found:")
        for err in invalid_statuses[:5]:
            print(f"   - {err}")
        return False

    print("✓ Schema & Score Validation: All ratings strictly in range [1, 5].")
    print("✓ Action Validation: All actions coherent with review status.")

    if require_complete and reviewed_count < 50:
        print(f"❌ Incomplete Review: require_complete=True but {pending_count} rows are still pending.")
        return False

    if reviewed_count == 50:
        print("✓ STATUS: COMPLETE (50/50 cases explicitly reviewed and validated)")
    else:
        print(f"✓ STATUS: WORKFLOW READY FOR REVIEW ({reviewed_count}/50 reviewed, {pending_count}/50 pending)")
    print("=" * 75)
    return True


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Validate human audit workbook integrity")
    parser.add_argument("--file", type=str, default=str(DEFAULT_XLSX), help="Path to audit workbook")
    parser.add_argument("--require-complete", action="store_true", help="Require all 50 rows to be reviewed")
    args = parser.parse_args()

    success = validate_review(Path(args.file), require_complete=args.require_complete)
    sys.exit(0 if success else 1)
