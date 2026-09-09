"""
Step 4F: Assisted Review of 50 Human-Audit Cases

Interactive and CLI workflow for reviewing the 50 human audit cases in:
eval/judge_human_review_ready_for_manual_check.xlsx

Commands:
  python scripts/review_human_audit.py status
  python scripts/review_human_audit.py show <message_id>
  python scripts/review_human_audit.py accept <message_id> [--notes "..."]
  python scripts/review_human_audit.py edit <message_id> --correctness C --helpfulness H --groundedness G --policy P --escalation E [--notes "..."]
  python scripts/review_human_audit.py review [--id <message_id>]
  python scripts/review_human_audit.py sync-csv
"""

import argparse
import csv
import sys
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple
import openpyxl

# UTF-8 stdout configuration
if sys.stdout.encoding != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

ROOT_DIR = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK_PATH = ROOT_DIR / "eval/judge_human_review_ready_for_manual_check.xlsx"
AUDIT_CSV_PATH = ROOT_DIR / "eval/judge_human_audit.csv"

# Columns in 'Judge-Human Audit'
COL_MESSAGE_ID = 1
COL_ORIGINAL_TEXT = 2
COL_AGENT_RESPONSE = 3
COL_RETRIEVED_EVIDENCE = 4
COL_AUDIT_FLAG = 5
COL_LLM_CORRECTNESS = 6
COL_ASST_CORRECTNESS = 7
COL_LLM_HELPFULNESS = 8
COL_ASST_HELPFULNESS = 9
COL_LLM_GROUNDEDNESS = 10
COL_ASST_GROUNDEDNESS = 11
COL_LLM_POLICY = 12
COL_ASST_POLICY = 13
COL_LLM_ESCALATION = 14
COL_ASST_ESCALATION = 15
COL_ASSISTANT_REASON = 16
COL_REVIEW_STATUS = 17
COL_HUMAN_CORRECTNESS = 18
COL_HUMAN_HELPFULNESS = 19
COL_HUMAN_GROUNDEDNESS = 20
COL_HUMAN_POLICY = 21
COL_HUMAN_ESCALATION = 22
COL_HUMAN_REVIEW_ACTION = 23
COL_HUMAN_NOTES = 24


def load_audit_sheet(workbook_path: Optional[Path] = None) -> Tuple[openpyxl.Workbook, Any]:
    target_path = workbook_path or DEFAULT_WORKBOOK_PATH
    if not target_path.exists():
        raise FileNotFoundError(f"Workbook not found: {target_path}")
    wb = openpyxl.load_workbook(target_path)
    if "Judge-Human Audit" not in wb.sheetnames:
        raise ValueError(f"Sheet 'Judge-Human Audit' not found in {target_path}")
    return wb, wb["Judge-Human Audit"]


def get_difficulty_highlights(ws: Any, r: int) -> List[str]:
    """
    Highlights difficult cases per Step 4F guidelines:
    - judge vs provisional rating differs by >= 2
    - escalation disagreement
    - policy < 4
    - groundedness < 4
    - low-confidence judge cases (scores == 3 or borderline)
    """
    reasons = []

    def _int(val: Any) -> Optional[int]:
        try:
            return int(val)
        except (ValueError, TypeError):
            return None

    llm_c = _int(ws.cell(row=r, column=COL_LLM_CORRECTNESS).value)
    llm_h = _int(ws.cell(row=r, column=COL_LLM_HELPFULNESS).value)
    llm_g = _int(ws.cell(row=r, column=COL_LLM_GROUNDEDNESS).value)
    llm_p = _int(ws.cell(row=r, column=COL_LLM_POLICY).value)
    llm_e = _int(ws.cell(row=r, column=COL_LLM_ESCALATION).value)

    hum_c = _int(ws.cell(row=r, column=COL_HUMAN_CORRECTNESS).value)
    hum_h = _int(ws.cell(row=r, column=COL_HUMAN_HELPFULNESS).value)
    hum_g = _int(ws.cell(row=r, column=COL_HUMAN_GROUNDEDNESS).value)
    hum_p = _int(ws.cell(row=r, column=COL_HUMAN_POLICY).value)
    hum_e = _int(ws.cell(row=r, column=COL_HUMAN_ESCALATION).value)

    dim_pairs = [
        ("Correctness", llm_c, hum_c),
        ("Helpfulness", llm_h, hum_h),
        ("Groundedness", llm_g, hum_g),
        ("Policy", llm_p, hum_p),
        ("Escalation", llm_e, hum_e),
    ]

    # 1. Rating differs by >= 2
    for name, l_val, h_val in dim_pairs:
        if l_val is not None and h_val is not None and abs(l_val - h_val) >= 2:
            reasons.append(f"{name} divergence >= 2 (Judge {l_val} vs Provisional {h_val})")

    # 2. Escalation disagreement
    if llm_e is not None and hum_e is not None and llm_e != hum_e:
        reasons.append(f"Escalation disagreement (Judge {llm_e} vs Provisional {hum_e})")

    # 3. Policy < 4
    if (hum_p is not None and hum_p < 4) or (llm_p is not None and llm_p < 4):
        reasons.append(f"Policy risk < 4 (Judge {llm_p}, Provisional {hum_p})")

    # 4. Groundedness < 4
    if (hum_g is not None and hum_g < 4) or (llm_g is not None and llm_g < 4):
        reasons.append(f"Groundedness gap < 4 (Judge {llm_g}, Provisional {hum_g})")

    # 5. Low-confidence judge cases
    judge_scores = [s for s in [llm_c, llm_h, llm_g, llm_p, llm_e] if s is not None]
    if any(s == 3 for s in judge_scores):
        reasons.append("Low-confidence judge case (borderline score of 3 present)")

    flag_col = str(ws.cell(row=r, column=COL_AUDIT_FLAG).value or "")
    if "Severe" in flag_col and not any("Severe" in r for r in reasons):
        reasons.append(flag_col)

    return reasons


def cmd_status(workbook_path: Optional[Path] = None):
    target_path = workbook_path or DEFAULT_WORKBOOK_PATH
    wb, ws = load_audit_sheet(target_path)
    total = 50
    reviewed = 0
    approved = 0
    edited = 0
    pending = 0
    difficult_cases = []

    for r in range(6, 56):
        mid = ws.cell(row=r, column=COL_MESSAGE_ID).value
        status = str(ws.cell(row=r, column=COL_REVIEW_STATUS).value or "").strip()
        action = str(ws.cell(row=r, column=COL_HUMAN_REVIEW_ACTION).value or "").strip()

        diff_highlights = get_difficulty_highlights(ws, r)
        if diff_highlights:
            difficult_cases.append((mid, diff_highlights))

        if status == "Reviewed" and action in ["approved", "edited"]:
            reviewed += 1
            if action == "approved":
                approved += 1
            elif action == "edited":
                edited += 1
        else:
            pending += 1

    progress_pct = (reviewed / total) * 100

    print("=" * 75)
    print("STEP 4F: ASSISTED HUMAN AUDIT STATUS DASHBOARD")
    print(f"Target: {target_path.name}")
    print("=" * 75)
    print(f"Total Audit Cases:     {total}")
    print(f"Reviewed:              {reviewed} / {total}")
    print(f"Pending:               {pending} / {total}")
    print(f"Progress:              {progress_pct:.1f}%")
    print(f"  - Approved Ratings:  {approved}")
    print(f"  - Edited Ratings:    {edited}")
    print(f"Difficult Cases Flagged: {len(difficult_cases)} / {total}")
    print("-" * 75)

    if pending > 0:
        print(f"Status: {pending} cases require explicit human review.")
        print("Run `python scripts/review_human_audit.py review` to inspect pending cases.")
    else:
        print("Status: All 50 cases have been explicitly reviewed.")
    print("=" * 75)


def cmd_show(mid: str, workbook_path: Optional[Path] = None):
    target_path = workbook_path or DEFAULT_WORKBOOK_PATH
    wb, ws = load_audit_sheet(target_path)
    target_row = None
    for r in range(6, 56):
        if str(ws.cell(row=r, column=COL_MESSAGE_ID).value).strip() == str(mid).strip():
            target_row = r
            break

    if not target_row:
        print(f"Error: Message ID {mid} not found in {target_path.name}.")
        return

    r = target_row
    status = str(ws.cell(row=r, column=COL_REVIEW_STATUS).value or "Needs Human Review").strip()
    action = str(ws.cell(row=r, column=COL_HUMAN_REVIEW_ACTION).value or "[None]").strip()
    notes = str(ws.cell(row=r, column=COL_HUMAN_NOTES).value or "").strip()

    highlights = get_difficulty_highlights(ws, r)

    print("=" * 75)
    print(f"AUDIT CASE: Message ID {mid}  |  Status: {status}  |  Action: {action}")
    print("=" * 75)

    if highlights:
        print("DIFFICULTY HIGHLIGHTS:")
        for h in highlights:
            print(f"  ⚠️  {h}")
        print("-" * 75)

    print(f"Customer Tweet:\n  {ws.cell(row=r, column=COL_ORIGINAL_TEXT).value}\n")
    print(f"Agent Response:\n  {ws.cell(row=r, column=COL_AGENT_RESPONSE).value}\n")
    print(f"Retrieved Evidence:\n  {ws.cell(row=r, column=COL_RETRIEVED_EVIDENCE).value}\n")
    print("-" * 75)
    print("SCORES COMPARISON:")
    print(f"  Dimension     | LLM Judge | Provisional Human Rating")
    print(f"  Correctness   |     {ws.cell(row=r, column=COL_LLM_CORRECTNESS).value}     |             {ws.cell(row=r, column=COL_HUMAN_CORRECTNESS).value}")
    print(f"  Helpfulness   |     {ws.cell(row=r, column=COL_LLM_HELPFULNESS).value}     |             {ws.cell(row=r, column=COL_HUMAN_HELPFULNESS).value}")
    print(f"  Groundedness  |     {ws.cell(row=r, column=COL_LLM_GROUNDEDNESS).value}     |             {ws.cell(row=r, column=COL_HUMAN_GROUNDEDNESS).value}")
    print(f"  Policy        |     {ws.cell(row=r, column=COL_LLM_POLICY).value}     |             {ws.cell(row=r, column=COL_HUMAN_POLICY).value}")
    print(f"  Escalation    |     {ws.cell(row=r, column=COL_LLM_ESCALATION).value}     |             {ws.cell(row=r, column=COL_HUMAN_ESCALATION).value}")
    print("-" * 75)
    print(f"Assistant Reason:\n  {ws.cell(row=r, column=COL_ASSISTANT_REASON).value}")
    if notes:
        print(f"Reviewer Notes:\n  {notes}")
    print("=" * 75)


def cmd_accept(mid: str, notes: str = "", workbook_path: Optional[Path] = None):
    target_path = workbook_path or DEFAULT_WORKBOOK_PATH
    wb, ws = load_audit_sheet(target_path)
    target_row = None
    for r in range(6, 56):
        if str(ws.cell(row=r, column=COL_MESSAGE_ID).value).strip() == str(mid).strip():
            target_row = r
            break

    if not target_row:
        print(f"Error: Message ID {mid} not found.")
        return

    r = target_row
    ratings = [
        int(ws.cell(row=r, column=COL_HUMAN_CORRECTNESS).value),
        int(ws.cell(row=r, column=COL_HUMAN_HELPFULNESS).value),
        int(ws.cell(row=r, column=COL_HUMAN_GROUNDEDNESS).value),
        int(ws.cell(row=r, column=COL_HUMAN_POLICY).value),
        int(ws.cell(row=r, column=COL_HUMAN_ESCALATION).value),
    ]

    ws.cell(row=r, column=COL_REVIEW_STATUS, value="Reviewed")
    ws.cell(row=r, column=COL_HUMAN_REVIEW_ACTION, value="approved")
    if notes:
        ws.cell(row=r, column=COL_HUMAN_NOTES, value=notes.strip())

    wb.save(target_path)
    print(f"✓ Message ID {mid}: Explicitly APPROVED provisional ratings {ratings}.")
    print(f"  review_status = 'Reviewed' | human_review_action = 'approved'")
    _print_quick_progress(ws)


def cmd_edit(mid: str, corr: int, helpf: int, ground: int, pol: int, esc: int, notes: str = "", workbook_path: Optional[Path] = None):
    for name, val in [("correctness", corr), ("helpfulness", helpf), ("groundedness", ground), ("policy", pol), ("escalation", esc)]:
        if val < 1 or val > 5:
            print(f"Error: {name} rating {val} must be an integer between 1 and 5.")
            return

    target_path = workbook_path or DEFAULT_WORKBOOK_PATH
    wb, ws = load_audit_sheet(target_path)
    target_row = None
    for r in range(6, 56):
        if str(ws.cell(row=r, column=COL_MESSAGE_ID).value).strip() == str(mid).strip():
            target_row = r
            break

    if not target_row:
        print(f"Error: Message ID {mid} not found.")
        return

    r = target_row
    ws.cell(row=r, column=COL_HUMAN_CORRECTNESS, value=int(corr))
    ws.cell(row=r, column=COL_HUMAN_HELPFULNESS, value=int(helpf))
    ws.cell(row=r, column=COL_HUMAN_GROUNDEDNESS, value=int(ground))
    ws.cell(row=r, column=COL_HUMAN_POLICY, value=int(pol))
    ws.cell(row=r, column=COL_HUMAN_ESCALATION, value=int(esc))
    ws.cell(row=r, column=COL_REVIEW_STATUS, value="Reviewed")
    ws.cell(row=r, column=COL_HUMAN_REVIEW_ACTION, value="edited")
    if notes:
        ws.cell(row=r, column=COL_HUMAN_NOTES, value=notes.strip())

    wb.save(target_path)
    print(f"✓ Message ID {mid}: Saved EDITED ratings [{corr}, {helpf}, {ground}, {pol}, {esc}].")
    print(f"  review_status = 'Reviewed' | human_review_action = 'edited'")
    _print_quick_progress(ws)


def _print_quick_progress(ws: Any):
    reviewed = 0
    total = 50
    for r in range(6, 56):
        status = str(ws.cell(row=r, column=COL_REVIEW_STATUS).value or "").strip()
        action = str(ws.cell(row=r, column=COL_HUMAN_REVIEW_ACTION).value or "").strip()
        if status == "Reviewed" and action in ["approved", "edited"]:
            reviewed += 1
    pending = total - reviewed
    pct = (reviewed / total) * 100
    print(f"Progress: Reviewed: {reviewed} / {total} | Pending: {pending} / {total} | {pct:.1f}%")


def cmd_interactive_review(specific_id: Optional[str] = None, workbook_path: Optional[Path] = None):
    target_path = workbook_path or DEFAULT_WORKBOOK_PATH
    wb, ws = load_audit_sheet(target_path)

    rows_to_review = []
    for r in range(6, 56):
        mid = str(ws.cell(row=r, column=COL_MESSAGE_ID).value).strip()
        status = str(ws.cell(row=r, column=COL_REVIEW_STATUS).value or "").strip()
        action = str(ws.cell(row=r, column=COL_HUMAN_REVIEW_ACTION).value or "").strip()

        if specific_id:
            if mid == str(specific_id).strip():
                rows_to_review.append(r)
                break
        else:
            if status != "Reviewed" or action not in ["approved", "edited"]:
                rows_to_review.append(r)

    if not rows_to_review:
        if specific_id:
            print(f"Message ID {specific_id} not found or already reviewed.")
        else:
            print("✓ All 50 audit cases have already been reviewed!")
        return

    print("=" * 75)
    print(f"STARTING ASSISTED REVIEW SESSION ({len(rows_to_review)} cases to review)")
    print("Options for each case:")
    print("  [A] Approve provisional ratings")
    print("  [E] Edit ratings")
    print("  [S] Skip to next")
    print("  [Q] Quit review session")
    print("=" * 75)

    for idx, r in enumerate(rows_to_review, start=1):
        mid = str(ws.cell(row=r, column=COL_MESSAGE_ID).value).strip()
        highlights = get_difficulty_highlights(ws, r)

        print("\n" + "=" * 75)
        print(f"CASE {idx}/{len(rows_to_review)}: Message ID {mid}")
        print("=" * 75)

        if highlights:
            print("⚠️  DIFFICULTY ALERTS:")
            for h in highlights:
                print(f"   • {h}")
            print("-" * 75)

        print(f"Customer Tweet:\n  {ws.cell(row=r, column=COL_ORIGINAL_TEXT).value}\n")
        print(f"Agent Response:\n  {ws.cell(row=r, column=COL_AGENT_RESPONSE).value}\n")
        print(f"Retrieved Evidence:\n  {ws.cell(row=r, column=COL_RETRIEVED_EVIDENCE).value}\n")
        print("-" * 75)

        cur_c = ws.cell(row=r, column=COL_HUMAN_CORRECTNESS).value
        cur_h = ws.cell(row=r, column=COL_HUMAN_HELPFULNESS).value
        cur_g = ws.cell(row=r, column=COL_HUMAN_GROUNDEDNESS).value
        cur_p = ws.cell(row=r, column=COL_HUMAN_POLICY).value
        cur_e = ws.cell(row=r, column=COL_HUMAN_ESCALATION).value

        print("SCORES COMPARISON:")
        print(f"  Dimension     | LLM Judge | Provisional Human Rating")
        print(f"  Correctness   |     {ws.cell(row=r, column=COL_LLM_CORRECTNESS).value}     |             {cur_c}")
        print(f"  Helpfulness   |     {ws.cell(row=r, column=COL_LLM_HELPFULNESS).value}     |             {cur_h}")
        print(f"  Groundedness  |     {ws.cell(row=r, column=COL_LLM_GROUNDEDNESS).value}     |             {cur_g}")
        print(f"  Policy        |     {ws.cell(row=r, column=COL_LLM_POLICY).value}     |             {cur_p}")
        print(f"  Escalation    |     {ws.cell(row=r, column=COL_LLM_ESCALATION).value}     |             {cur_e}")
        print("-" * 75)
        print(f"Assistant Reason:\n  {ws.cell(row=r, column=COL_ASSISTANT_REASON).value}")
        print("-" * 75)

        while True:
            choice = input("Choice: [A]pprove provisional, [E]dit, [S]kip, [Q]uit: ").strip().upper()
            if choice == "A":
                notes = input("Optional notes (press Enter to leave empty): ").strip()
                ws.cell(row=r, column=COL_REVIEW_STATUS, value="Reviewed")
                ws.cell(row=r, column=COL_HUMAN_REVIEW_ACTION, value="approved")
                if notes:
                    ws.cell(row=r, column=COL_HUMAN_NOTES, value=notes)
                wb.save(target_path)
                print(f"✓ Case {mid} APPROVED.")
                _print_quick_progress(ws)
                break
            elif choice == "E":
                def _prompt_score(name: str, default: Any) -> int:
                    while True:
                        inp = input(f"  {name} [1-5] (default {default}): ").strip()
                        if not inp:
                            return int(default)
                        try:
                            val = int(inp)
                            if 1 <= val <= 5:
                                return val
                            print("  Must be between 1 and 5.")
                        except ValueError:
                            print("  Must be an integer.")

                new_c = _prompt_score("Correctness", cur_c)
                new_h = _prompt_score("Helpfulness", cur_h)
                new_g = _prompt_score("Groundedness", cur_g)
                new_p = _prompt_score("Policy", cur_p)
                new_e = _prompt_score("Escalation", cur_e)
                notes = input("  Optional notes (press Enter to leave empty): ").strip()

                ws.cell(row=r, column=COL_HUMAN_CORRECTNESS, value=new_c)
                ws.cell(row=r, column=COL_HUMAN_HELPFULNESS, value=new_h)
                ws.cell(row=r, column=COL_HUMAN_GROUNDEDNESS, value=new_g)
                ws.cell(row=r, column=COL_HUMAN_POLICY, value=new_p)
                ws.cell(row=r, column=COL_HUMAN_ESCALATION, value=new_e)
                ws.cell(row=r, column=COL_REVIEW_STATUS, value="Reviewed")
                ws.cell(row=r, column=COL_HUMAN_REVIEW_ACTION, value="edited")
                if notes:
                    ws.cell(row=r, column=COL_HUMAN_NOTES, value=notes)
                wb.save(target_path)
                print(f"✓ Case {mid} EDITED with [{new_c}, {new_h}, {new_g}, {new_p}, {new_e}].")
                _print_quick_progress(ws)
                break
            elif choice == "S":
                print(f"Skipping case {mid}.")
                break
            elif choice == "Q":
                print("Exiting review session.")
                return
            else:
                print("Invalid choice. Please select A, E, S, or Q.")


def cmd_sync_csv(workbook_path: Optional[Path] = None):
    target_path = workbook_path or DEFAULT_WORKBOOK_PATH
    wb, ws = load_audit_sheet(target_path)
    if not AUDIT_CSV_PATH.exists():
        print(f"Error: {AUDIT_CSV_PATH} not found.")
        return

    with open(AUDIT_CSV_PATH, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        fieldnames = reader.fieldnames
        rows = list(reader)

    sheet_by_id = {}
    for r in range(6, 56):
        mid = str(ws.cell(row=r, column=COL_MESSAGE_ID).value).strip()
        status = str(ws.cell(row=r, column=COL_REVIEW_STATUS).value or "").strip()
        action = str(ws.cell(row=r, column=COL_HUMAN_REVIEW_ACTION).value or "").strip()
        ratings = [ws.cell(row=r, column=c).value for c in range(COL_HUMAN_CORRECTNESS, COL_HUMAN_ESCALATION + 1)]
        notes = str(ws.cell(row=r, column=COL_HUMAN_NOTES).value or "")
        sheet_by_id[mid] = (status, action, ratings, notes)

    updated = 0
    for r in rows:
        mid = str(r["message_id"]).strip()
        if mid in sheet_by_id:
            status, action, ratings, notes = sheet_by_id[mid]
            if status == "Reviewed" and action in ["approved", "edited"]:
                r["human_correctness"] = str(ratings[0])
                r["human_helpfulness"] = str(ratings[1])
                r["human_groundedness"] = str(ratings[2])
                r["human_policy"] = str(ratings[3])
                r["human_escalation"] = str(ratings[4])
                r["human_notes"] = notes
                updated += 1

    with open(AUDIT_CSV_PATH, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    print(f"✓ Synchronized {updated} human review ratings to {AUDIT_CSV_PATH.name}.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Step 4F: Assisted Human Audit Review Tool")
    parser.add_argument("--file", type=str, default=str(DEFAULT_WORKBOOK_PATH), help="Workbook path")
    sub = parser.add_subparsers(dest="subcommand")

    sub.add_parser("status")

    p_show = sub.add_parser("show")
    p_show.add_argument("message_id", type=str)

    p_accept = sub.add_parser("accept")
    p_accept.add_argument("message_id", type=str)
    p_accept.add_argument("--notes", type=str, default="")

    p_edit = sub.add_parser("edit")
    p_edit.add_argument("message_id", type=str)
    p_edit.add_argument("--correctness", type=int, required=True)
    p_edit.add_argument("--helpfulness", type=int, required=True)
    p_edit.add_argument("--groundedness", type=int, required=True)
    p_edit.add_argument("--policy", type=int, required=True)
    p_edit.add_argument("--escalation", type=int, required=True)
    p_edit.add_argument("--notes", type=str, default="")

    p_rev = sub.add_parser("review")
    p_rev.add_argument("--id", type=str, default=None, help="Optional specific message ID to review")

    sub.add_parser("sync-csv")

    args = parser.parse_args()
    wb_file = Path(args.file)

    if args.subcommand == "status":
        cmd_status(wb_file)
    elif args.subcommand == "show":
        cmd_show(args.message_id, wb_file)
    elif args.subcommand == "accept":
        cmd_accept(args.message_id, args.notes, wb_file)
    elif args.subcommand == "edit":
        cmd_edit(args.message_id, args.correctness, args.helpfulness, args.groundedness, args.policy, args.escalation, args.notes, wb_file)
    elif args.subcommand == "review":
        cmd_interactive_review(args.id, wb_file)
    elif args.subcommand == "sync-csv":
        cmd_sync_csv(wb_file)
    else:
        cmd_status(wb_file)
