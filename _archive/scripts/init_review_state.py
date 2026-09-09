"""
Step 2C: Initialize Human Review State and Conversation Thread Cache

Extracts:
1. eval/golden_threads_cache.json (All conversation threads for the 200 golden examples)
2. eval/golden_review_state.json (200 review items with assistant proposals, difficult case tags, and strictly empty human fields)
"""

import csv
import json
import os
import sys
from pathlib import Path
import openpyxl

# Reconfigure stdout for UTF-8
if sys.stdout.encoding != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

EVAL_DIR = Path("eval")
DATA_DIR = Path("data/processed")

CSV_PATH = EVAL_DIR / "golden_eval_set.csv"
ASSISTED_XLSX_PATH = EVAL_DIR / "golden_review_assisted.xlsx"
THREADS_JSONL_PATH = DATA_DIR / "AmazonHelp_threads.jsonl"
EXCLUSIONS_JSON_PATH = EVAL_DIR / "golden_thread_exclusions.json"

THREADS_CACHE_PATH = EVAL_DIR / "golden_threads_cache.json"
REVIEW_STATE_PATH = EVAL_DIR / "golden_review_state.json"

LOCKED_INTENTS = [
    "Delivery Tracking & Status",
    "Delivery Problem & Logistics",
    "Returns, Replacements & Refunds",
    "Payment, Billing & Gift Cards",
    "Prime & Subscription Services",
    "Order & Checkout",
    "Account Access & Security",
    "Digital Services & Devices",
    "Seller & Product Quality",
    "General / Feedback / Other",
]

ALLOWED_LANGUAGES = ["en", "es", "ja", "de", "pt", "fr", "it", "other"]

ALLOWED_STATES = [
    "new_issue",
    "active_troubleshooting",
    "dm_handoff",
    "follow_up",
    "resolved_or_acknowledgment",
    "unclear",
]

BOUNDARY_PAIRS = [
    {"Delivery Tracking & Status", "Delivery Problem & Logistics"},
    {"Prime & Subscription Services", "Delivery Problem & Logistics"},
    {"Returns, Replacements & Refunds", "Payment, Billing & Gift Cards"},
    {"Returns, Replacements & Refunds", "Seller & Product Quality"},
    {"Account Access & Security", "General / Feedback / Other"},
]


def extract_thread_cache(target_conv_ids: set) -> dict:
    print(f"Extracting thread turns for {len(target_conv_ids)} conversations from {THREADS_JSONL_PATH}...")
    threads_cache = {}
    with open(THREADS_JSONL_PATH, "r", encoding="utf-8") as f:
        for line in f:
            data = json.loads(line)
            cid = str(data.get("conversation_id"))
            if cid in target_conv_ids:
                threads_cache[cid] = {
                    "conversation_id": cid,
                    "total_turns": data.get("total_turns", len(data.get("turns", []))),
                    "turns": data.get("turns", [])
                }
                if len(threads_cache) == len(target_conv_ids):
                    break
    print(f"✓ Cached {len(threads_cache)} threads.")
    return threads_cache


def init_review_state():
    print("=" * 75)
    print("STEP 2C: INITIALIZING HUMAN REVIEW STATE & THREAD CACHE")
    print("=" * 75)

    # 1. Load golden eval set CSV to get conversation_ids and authentic text
    if not CSV_PATH.exists():
        raise FileNotFoundError(f"Missing {CSV_PATH}")
    with open(CSV_PATH, "r", encoding="utf-8") as f:
        csv_rows = list(csv.DictReader(f))

    mid_to_conv = {int(r["message_id"]): r["conversation_id"] for r in csv_rows}
    mid_to_text = {int(r["message_id"]): r["original_text"] for r in csv_rows}
    target_conv_ids = set(mid_to_conv.values())

    # 2. Extract and cache conversation threads
    threads_cache = extract_thread_cache(target_conv_ids)
    with open(THREADS_CACHE_PATH, "w", encoding="utf-8") as f:
        json.dump(threads_cache, f, indent=2, ensure_ascii=False)
    print(f"✓ Saved thread cache to {THREADS_CACHE_PATH}")

    # 3. Load assistant proposals from golden_review_assisted.xlsx
    if not ASSISTED_XLSX_PATH.exists():
        raise FileNotFoundError(f"Missing {ASSISTED_XLSX_PATH}")

    wb = openpyxl.load_workbook(ASSISTED_XLSX_PATH)
    ws = wb["Review"]

    items = []
    difficult_count = 0

    for r_idx in range(6, 206):
        mid = int(ws.cell(r_idx, 1).value)
        text = mid_to_text.get(mid, ws.cell(r_idx, 2).value)
        conv_id = mid_to_conv.get(mid, "")

        prop_intent = str(ws.cell(r_idx, 3).value or "").strip()
        prop_lang = str(ws.cell(r_idx, 5).value or "en").strip()
        prop_state = str(ws.cell(r_idx, 7).value or "new_issue").strip()
        prop_esc = str(ws.cell(r_idx, 9).value).lower() in ["true", "1", "yes"]
        prop_sec = str(ws.cell(r_idx, 11).value).lower() in ["true", "1", "yes"]
        prop_prio = str(ws.cell(r_idx, 13).value or "standard").strip()
        prop_reason = str(ws.cell(r_idx, 15).value or "").strip()
        prop_conf = str(ws.cell(r_idx, 16).value or "MEDIUM").strip().upper()
        ambiguity_note = str(ws.cell(r_idx, 17).value or "").strip()
        if ambiguity_note == "None":
            ambiguity_note = ""

        # Difficult case criteria:
        # 1. Low confidence
        # 2. Ambiguity note exists
        # 3. Security alert = true
        # 4. Escalation = true
        # 5. Multilingual (lang != 'en')
        # 6. Proposed intent boundary pairs
        difficult_reasons = []
        if prop_conf == "LOW":
            difficult_reasons.append("Low Assistant Confidence (LOW)")
        if ambiguity_note:
            difficult_reasons.append(f"Ambiguity Note: {ambiguity_note}")
        if prop_sec:
            difficult_reasons.append("Proposed Security Alert (P0_CRITICAL)")
        if prop_esc:
            difficult_reasons.append("Escalation Flagged")
        if prop_lang != "en":
            difficult_reasons.append(f"Multilingual Message ({prop_lang})")

        for pair in BOUNDARY_PAIRS:
            if prop_intent in pair:
                other = list(pair - {prop_intent})[0]
                if ambiguity_note and other.lower() in ambiguity_note.lower():
                    difficult_reasons.append(f"Boundary Conflict: {prop_intent} vs {other}")

        is_difficult = len(difficult_reasons) > 0
        if is_difficult:
            difficult_count += 1

        item = {
            "index": r_idx - 5,
            "message_id": mid,
            "conversation_id": conv_id,
            "original_text": text,
            "assistant_proposed_intent": prop_intent,
            "assistant_proposed_language": prop_lang,
            "assistant_proposed_conversation_state": prop_state,
            "assistant_proposed_escalate": prop_esc,
            "assistant_proposed_is_security_alert": prop_sec,
            "assistant_proposed_priority": prop_prio,
            "assistant_proposed_reason": prop_reason,
            "assistant_confidence": prop_conf,
            "assistant_ambiguity_note": ambiguity_note,
            "is_difficult": is_difficult,
            "difficult_reasons": difficult_reasons,
            # Human fields - strictly null/empty initially
            "my_final_intent": None,
            "my_final_language": None,
            "my_final_conversation_state": None,
            "my_final_escalate": None,
            "my_final_is_security_alert": None,
            "my_final_priority": None,
            "notes": "",
            "human_review_action": None,  # 'approved' or 'edited'
            "is_reviewed": False,
        }
        items.append(item)

    review_state = {
        "metadata": {
            "total_examples": len(items),
            "reviewed_count": 0,
            "pending_count": len(items),
            "difficult_count": difficult_count,
            "approved_count": 0,
            "edited_count": 0,
        },
        "items": items,
    }

    with open(REVIEW_STATE_PATH, "w", encoding="utf-8") as f:
        json.dump(review_state, f, indent=2, ensure_ascii=False)

    print(f"✓ Initialized review state at {REVIEW_STATE_PATH}")
    print(f"  - Total items:           {len(items)}")
    print(f"  - Pending review:        {len(items)}")
    print(f"  - Flagged difficult:     {difficult_count}")
    print(f"  - Human fields:          ALL EMPTY (No automated fabrication)")


if __name__ == "__main__":
    init_review_state()
