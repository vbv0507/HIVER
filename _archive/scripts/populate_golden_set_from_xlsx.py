"""
Step 2: Populate eval/golden_eval_set.csv from eval/golden_eval_set_assistant_filled.xlsx

Transfers the completed review annotations from the assistant-filled Excel file
into the canonical CSV format while preserving data integrity and encoding.
"""

import csv
import sys
from pathlib import Path
import openpyxl

# Reconfigure stdout for UTF-8
if sys.stdout.encoding != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

EVAL_DIR = Path("eval")
XLSX_PATH = EVAL_DIR / "golden_eval_set_assistant_filled.xlsx"
CSV_PATH = EVAL_DIR / "golden_eval_set.csv"
BACKUP_PATH = EVAL_DIR / "golden_eval_set_pre_review_backup.csv"

EXPECTED_HEADERS = [
    "message_id",
    "conversation_id",
    "original_text",
    "ai_suggested_intent",
    "my_final_intent",
    "ai_suggested_language",
    "my_final_language",
    "ai_suggested_conversation_state",
    "my_final_conversation_state",
    "ai_suggested_escalate",
    "my_final_escalate",
    "ai_suggested_reason",
    "is_security_alert",
    "my_final_is_security_alert",
    "priority",
    "my_final_priority",
    "notes",
]


def populate_golden_set() -> None:
    if not XLSX_PATH.exists():
        raise FileNotFoundError(f"Source Excel workbook not found: {XLSX_PATH}")

    # Ensure pre-review backup exists
    if not BACKUP_PATH.exists() and CSV_PATH.exists():
        import shutil
        shutil.copyfile(CSV_PATH, BACKUP_PATH)
        print(f"Created pre-review backup at {BACKUP_PATH}")

    print(f"Loading {XLSX_PATH}...")
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb["Review"] if "Review" in wb.sheetnames else wb.active

    xlsx_rows = list(ws.iter_rows(values_only=True))
    xlsx_header = [str(c).strip() if c is not None else "" for c in xlsx_rows[0]]
    data_rows = xlsx_rows[1:]

    print(f"Loaded {len(data_rows)} data rows from sheet '{ws.title}'.")

    # Map column positions by header name
    col_map = {name: idx for idx, name in enumerate(xlsx_header)}
    for h in EXPECTED_HEADERS:
        if h not in col_map:
            raise ValueError(f"Expected column '{h}' not found in Excel headers: {xlsx_header}")

    # Process rows
    processed_rows = []
    for r_idx, r in enumerate(data_rows, start=2):
        row_dict = {}
        for h in EXPECTED_HEADERS:
            val = r[col_map[h]]
            if val is None:
                val = ""
            elif isinstance(val, bool):
                val = "true" if val else "false"
            else:
                val_str = str(val).strip()
                # Normalize boolean strings for escalate and security alert
                if h in ["ai_suggested_escalate", "my_final_escalate", "is_security_alert", "my_final_is_security_alert"]:
                    if val_str.lower() in ["true", "1", "yes"]:
                        val = "true"
                    elif val_str.lower() in ["false", "0", "no"]:
                        val = "false"
                    else:
                        val = val_str
                else:
                    val = val_str
            row_dict[h] = val
        processed_rows.append(row_dict)

    print(f"Writing {len(processed_rows)} rows to {CSV_PATH}...")
    with open(CSV_PATH, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=EXPECTED_HEADERS)
        writer.writeheader()
        writer.writerows(processed_rows)

    print("✓ eval/golden_eval_set.csv updated successfully.")


if __name__ == "__main__":
    populate_golden_set()
