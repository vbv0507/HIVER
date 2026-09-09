"""
Step 4E: Pre-fill 50 Audit Ratings for Human Review

Produces eval/judge_human_review_ready_for_manual_check.xlsx:
- Pre-fills the five human rating columns with provisional assistant ratings
  so the human reviewer can inspect and calibrate each row manually.
- Adds visible banner:
  "These pre-filled human ratings are provisional and require independent review."
- Adds review_status column with initial value: "Needs Human Review".
- Does NOT create fake "approved" or "edited" audit actions.
- Does NOT fabricate human authorship or notes.
- Strictly preserves LLM judge machine values and the Assistant Suggestions sheet.
"""

import sys
from pathlib import Path
from typing import Dict, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

# Reconfigure stdout for UTF-8 compatibility
if sys.stdout.encoding != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

ROOT_DIR = Path(__file__).resolve().parent.parent
SRC_XLSX_1 = ROOT_DIR / "eval/judge_human_review_prefilled_for_review.xlsx"
SRC_XLSX_2 = ROOT_DIR / "eval/judge_human_review_ready.xlsx"
OUT_XLSX = ROOT_DIR / "eval/judge_human_review_ready_for_manual_check.xlsx"

# Exact 50 provisional ratings:
# message_id -> (correctness, helpfulness, groundedness, policy, escalation)
PROVISIONAL_RATINGS: Dict[str, Tuple[int, int, int, int, int]] = {
    "25279": (2, 2, 2, 5, 2),
    "65900": (5, 5, 4, 5, 5),
    "77408": (1, 1, 2, 5, 4),
    "77535": (2, 2, 2, 5, 3),
    "95437": (5, 5, 3, 5, 5),
    "179265": (1, 1, 2, 5, 2),
    "180800": (1, 1, 2, 5, 2),
    "183483": (4, 3, 5, 5, 5),
    "234555": (1, 1, 2, 5, 2),
    "258713": (1, 1, 2, 5, 2),
    "291347": (1, 1, 2, 5, 2),
    "379700": (1, 1, 2, 5, 1),
    "379788": (1, 1, 2, 5, 2),
    "459028": (5, 5, 4, 5, 5),
    "483090": (4, 3, 4, 5, 5),
    "490648": (4, 3, 4, 5, 5),
    "498538": (1, 1, 1, 5, 1),
    "508166": (1, 1, 2, 5, 2),
    "532639": (1, 1, 1, 5, 1),
    "649952": (1, 1, 1, 5, 1),
    "679078": (1, 1, 1, 5, 1),
    "701334": (1, 1, 1, 5, 1),
    "909507": (5, 5, 4, 5, 5),
    "944251": (1, 1, 2, 5, 1),
    "947686": (1, 1, 2, 5, 2),
    "984200": (5, 4, 3, 5, 3),
    "1113514": (1, 1, 1, 5, 1),
    "1300626": (1, 1, 2, 5, 2),
    "1313039": (1, 1, 2, 5, 1),
    "1362458": (1, 1, 2, 5, 2),
    "1382396": (1, 1, 1, 5, 1),
    "1572905": (1, 1, 2, 5, 2),
    "1691622": (1, 1, 1, 5, 2),
    "1764531": (5, 4, 3, 5, 5),
    "1829432": (1, 1, 2, 5, 2),
    "2126759": (4, 3, 3, 5, 2),
    "2184761": (1, 1, 1, 5, 2),
    "2254381": (1, 1, 1, 5, 2),
    "2465040": (1, 1, 1, 5, 1),
    "2531101": (2, 1, 1, 5, 2),
    "2541046": (2, 2, 3, 5, 3),
    "2611666": (2, 2, 2, 5, 2),
    "2777430": (2, 2, 3, 5, 2),
    "2803279": (1, 1, 1, 5, 3),
    "2841828": (1, 1, 2, 5, 1),
    "2899522": (1, 1, 1, 5, 1),
    "2922515": (1, 1, 1, 5, 1),
    "2926332": (4, 3, 2, 5, 5),
    "2927090": (3, 3, 3, 5, 4),
    "2971945": (3, 2, 2, 5, 4),
}


def prefill_workbook():
    print("=" * 75)
    print("STEP 4E: PRE-FILLING PROVISIONAL RATINGS FOR HUMAN AUDIT")
    print("=" * 75)

    src_file = SRC_XLSX_1 if SRC_XLSX_1.exists() else SRC_XLSX_2
    if not src_file.exists():
        raise FileNotFoundError(f"Source file not found: {SRC_XLSX_1} or {SRC_XLSX_2}")

    print(f"Loading from: {src_file}")
    wb = openpyxl.load_workbook(src_file)

    if "Judge-Human Audit" not in wb.sheetnames:
        raise ValueError("Sheet 'Judge-Human Audit' not found in source workbook.")
    if "Assistant Suggestions" not in wb.sheetnames:
        raise ValueError("Sheet 'Assistant Suggestions' not found in source workbook.")

    ws = wb["Judge-Human Audit"]

    # 1. Add / Update Visible Banner in Row 1
    banner_text = "NOTICE: These pre-filled human ratings are provisional and require independent review."
    ws.merge_cells("A1:Y1")
    banner_cell = ws.cell(row=1, column=1, value=banner_text)
    banner_cell.font = Font(name="Calibri", size=11, bold=True, color="991B1B")
    banner_cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # Amber-100
    banner_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # 2. Row 2 Dashboard Update
    ws.cell(row=2, column=1, value="AUDIT STATUS:").font = Font(name="Calibri", size=10, bold=True, color="1E3A8A")
    ws.cell(row=2, column=2, value="50 Provisional Drafts").font = Font(name="Calibri", size=10, bold=True, color="D97706")
    ws.cell(row=2, column=3, value="Review State:").font = Font(name="Calibri", size=10, bold=True, color="991B1B")
    ws.cell(row=2, column=4, value="Needs Human Review (50)").font = Font(name="Calibri", size=10, bold=True, color="991B1B")
    ws.cell(row=2, column=5, value="Machine Values:").font = Font(name="Calibri", size=10, bold=True, color="059669")
    ws.cell(row=2, column=6, value="100% Pristine").font = Font(name="Calibri", size=10, bold=True, color="059669")
    ws.row_dimensions[2].height = 22

    # 3. Row 3 Instructions Update
    ws.cell(
        row=3, column=1,
        value="INSTRUCTIONS FOR HUMAN REVIEWER: Inspect each row. If you agree with the provisional rating, leave it. If you disagree, change the score (1–5). Update review_status when finished."
    ).font = Font(name="Calibri", size=9, italic=True, color="1E293B")
    ws.row_dimensions[3].height = 22

    # 4. Super-Header Row 4
    ws.merge_cells("Q4:X4")
    sh_cell = ws.cell(row=4, column=17, value="3. PROVISIONAL HUMAN AUDIT RATINGS (REQUIRES MANUAL CHECK)")
    sh_cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    sh_cell.fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")  # Teal-600
    sh_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 5. Column Headers in Row 5
    # Col 17: review_status
    # Col 18: human_correctness
    # Col 19: human_helpfulness
    # Col 20: human_groundedness
    # Col 21: human_policy
    # Col 22: human_escalation
    # Col 23: human_review_action (blank)
    # Col 24: human_notes (blank)
    fill_hdr_status = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")  # Amber-600
    fill_hdr_human = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")   # Teal-600
    thin_border_side = Side(style="thin", color="CBD5E1")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    ws.cell(row=5, column=17, value="review_status").font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    ws.cell(row=5, column=17).fill = fill_hdr_status
    ws.cell(row=5, column=17).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=5, column=17).border = border_cell

    for c in range(18, 23):
        ws.cell(row=5, column=c).fill = fill_hdr_human

    ws.cell(row=5, column=23, value="human_review_action").font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    ws.cell(row=5, column=23).fill = fill_hdr_human
    ws.cell(row=5, column=24, value="human_notes").font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    ws.cell(row=5, column=24).fill = fill_hdr_human

    # 6. Populate 50 rows (rows 6 to 55)
    data_start = 6
    populated_count = 0

    font_provisional = Font(name="Calibri", size=10, bold=True, color="0F766E")  # Teal-700
    fill_provisional = PatternFill(start_color="F0FDFA", end_color="F0FDFA", fill_type="solid")  # Teal-50
    font_status = Font(name="Calibri", size=9, bold=True, color="B45309")  # Amber-700
    fill_status = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    for r in range(data_start, data_start + 50):
        mid = str(ws.cell(row=r, column=1).value).strip()
        if mid not in PROVISIONAL_RATINGS:
            raise KeyError(f"Row {r}: message_id '{mid}' not found in provisional ratings!")

        ratings = PROVISIONAL_RATINGS[mid]

        # Col 17: review_status
        c_status = ws.cell(row=r, column=17, value="Needs Human Review")
        c_status.font = font_status
        c_status.fill = fill_status
        c_status.alignment = align_center
        c_status.border = border_cell

        # Col 18-22: 5 human rating fields
        for dim_idx, score in enumerate(ratings):
            col_num = 18 + dim_idx
            cell = ws.cell(row=r, column=col_num, value=int(score))
            cell.font = font_provisional
            cell.fill = fill_provisional
            cell.alignment = align_center
            cell.border = border_cell

        # Col 23: human_review_action (leave blank - no fake approved/edited audit actions)
        c_action = ws.cell(row=r, column=23, value="")
        c_action.font = Font(name="Calibri", size=9)
        c_action.alignment = align_center
        c_action.border = border_cell

        # Col 24: human_notes (leave blank - no invented notes)
        c_notes = ws.cell(row=r, column=24, value="")
        c_notes.font = Font(name="Calibri", size=9)
        c_notes.alignment = align_left
        c_notes.border = border_cell

        populated_count += 1

    # Width adjustment
    ws.column_dimensions["Q"].width = 22
    for col_letter in ["R", "S", "T", "U", "V"]:
        ws.column_dimensions[col_letter].width = 17
    ws.column_dimensions["W"].width = 20
    ws.column_dimensions["X"].width = 25

    # 7. Preserve Assistant Suggestions sheet
    # Check that Assistant Suggestions exists and has 50 rows
    ws_asst = wb["Assistant Suggestions"]
    if ws_asst.max_row < 50:
        raise ValueError("Assistant Suggestions sheet is incomplete!")

    # 8. Save output workbook
    wb.save(OUT_XLSX)
    print(f"✓ Saved pre-filled workbook to: {OUT_XLSX}")
    print(f"✓ Populated {populated_count} rows with provisional ratings.")
    print("✓ Added visible banner: 'These pre-filled human ratings are provisional and require independent review.'")
    print("✓ Added column 'review_status' with value 'Needs Human Review'.")
    print("✓ Zero fake 'approved' or 'edited' audit actions created.")
    print("✓ Preserved all LLM judge scores and Assistant Suggestions sheet.")


if __name__ == "__main__":
    prefill_workbook()
